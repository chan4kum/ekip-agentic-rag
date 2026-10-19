"""Build small, valid sample documents in memory (no binary fixtures in the repo)."""

import base64
import io
import os
import zipfile

from docx import Document as new_docx
from PIL import Image


def png(width: int = 160, height: int = 120, *, noise: bool = True) -> bytes:
    """A PNG. Random noise keeps it above the 'meaningful image' byte threshold."""
    raw = os.urandom(width * height * 3) if noise else bytes(width * height * 3)
    img = Image.frombytes("RGB", (width, height), raw)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def data_uri(data: bytes, mime: str = "image/png") -> str:
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


def docx(
    *,
    headings: bool = True,
    table: bool = True,
    picture: bytes | None = None,
    extra_bytes: int = 0,
) -> bytes:
    d = new_docx()
    d.core_properties.title = "Maintenance Manual"
    if headings:
        d.add_heading("Chapter 1", level=1)
        d.add_heading("Torque values", level=2)
    d.add_paragraph("Tighten the bolts in a star pattern.")
    d.add_paragraph("Check the gasket", style="List Bullet")
    if table:
        t = d.add_table(rows=2, cols=2)
        for r, row in enumerate([["Bolt", "Torque"], ["M8", "25 Nm"]]):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    if picture is not None:
        d.add_picture(io.BytesIO(picture))
    buf = io.BytesIO()
    d.save(buf)
    if extra_bytes:  # add an uncompressed filler entry to inflate the zip's declared size
        with zipfile.ZipFile(buf, "a", zipfile.ZIP_STORED) as zf:
            zf.writestr("filler.bin", bytes(extra_bytes))
    return buf.getvalue()


def pdf_with_text(pages: list[list[str]]) -> bytes:
    """A minimal valid text PDF: each page is a list of text lines (Helvetica 12)."""
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        page_id, content_id = 4 + 2 * i, 5 + 2 * i  # noqa: F841 - documents the numbering
        stream_lines = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_lines.append(f"({safe}) Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def scanned_pdf(images: list[bytes]) -> bytes:
    """A PDF whose pages are just pictures (no text layer), like a scan."""
    pil = [Image.open(io.BytesIO(i)).convert("RGB") for i in images]
    buf = io.BytesIO()
    pil[0].save(buf, "PDF", save_all=True, append_images=pil[1:])
    return buf.getvalue()
